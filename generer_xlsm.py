# -*- coding: utf-8 -*-
"""
generer_xlsm.py
================================================================
Génère le classeur professionnel :
    Tableau_suivi_gasoil_et_production_eau_v2.xlsm

Fonctionnement :
1. openpyxl crée un fichier .xlsx valide.
2. Les modules VBA sont générés dans le dossier ./vba/.
3. Si pywin32 + Excel sont disponibles, le script :
   - ouvre le .xlsx ;
   - importe les modules VBA ;
   - enregistre le classeur en .xlsm (FileFormat = 52).
4. Si l'injection échoue, le .xlsx reste utilisable et les
   instructions d'import manuel sont affichées.

Prérequis :
    pip install openpyxl
    pip install pywin32   # optionnel, Windows + Excel uniquement
================================================================
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# ----------------------------------------------------------------
# CONFIGURATION GÉNÉRALE
# ----------------------------------------------------------------

FICHIER_XLSX = "Tableau_suivi_gasoil_et_production_eau_v2.xlsx"
FICHIER_XLSM = "Tableau_suivi_gasoil_et_production_eau_v2.xlsm"
DOSSIER_VBA = "vba"

# Palette de couleurs
BLEU_FONCE = "1F4E78"
BLEU_MOYEN = "2E75B6"
BLEU_CLAIR = "D9E1F2"
JAUNE_SAISIE = "FFF2CC"
GRIS_CALCUL = "F2F2F2"
ROUGE = "FFC7CE"
ORANGE = "FFEB9C"
VERT = "C6EFCE"
GRIS_BORDURE = "BFBFBF"


# ----------------------------------------------------------------
# HELPERS DE STYLE
# ----------------------------------------------------------------

def fill(hex_color):
    """Retourne un remplissage uni."""
    return PatternFill("solid", fgColor=hex_color)


def bordure_fine():
    """Retourne une bordure fine standard."""
    side = Side(style="thin", color=GRIS_BORDURE)
    return Border(left=side, right=side, top=side, bottom=side)


def style_titre(ws, cell_range, texte, taille=16):
    """Applique un titre principal bleu foncé."""
    ws.merge_cells(cell_range)
    first_cell = cell_range.split(":")[0]
    c = ws[first_cell]
    c.value = texte
    c.font = Font(bold=True, color="FFFFFF", size=taille)
    c.fill = fill(BLEU_FONCE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[c.row].height = 28


def style_section(ws, cell_range, texte):
    """Applique un bandeau de section bleu clair."""
    ws.merge_cells(cell_range)
    first_cell = cell_range.split(":")[0]
    c = ws[first_cell]
    c.value = texte
    c.font = Font(bold=True, color=BLEU_FONCE)
    c.fill = fill(BLEU_CLAIR)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[c.row].height = 22


def style_entetes(ws, ligne, colonnes):
    """Met en forme une ligne d'en-têtes."""
    for col in colonnes:
        c = ws.cell(row=ligne, column=col)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = fill(BLEU_MOYEN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bordure_fine()
    ws.row_dimensions[ligne].height = 34


def style_cellule(ws, coord, type_zone="saisie", alignement="left"):
    """Applique le style saisie / calcul / normal."""
    c = ws[coord]
    c.border = bordure_fine()

    if type_zone == "saisie":
        c.fill = fill(JAUNE_SAISIE)
    elif type_zone == "calcul":
        c.fill = fill(GRIS_CALCUL)

    c.alignment = Alignment(
        horizontal=alignement,
        vertical="center",
        wrap_text=True
    )


def definir_largeurs(ws, largeurs):
    """Définit les largeurs de colonnes : {'A': 12, 'B': 20}."""
    for col, largeur in largeurs.items():
        ws.column_dimensions[col].width = largeur


def ajouter_validation_liste(ws, formule, plage):
    """Ajoute une validation de type liste."""
    dv = DataValidation(type="list", formula1=formule, allow_blank=True)
    dv.error = "Valeur non autorisée. Sélectionnez une valeur de la liste."
    dv.errorTitle = "Validation"
    ws.add_data_validation(dv)
    dv.add(plage)
    return dv


def ajouter_validation_decimal_positif(ws, plage):
    """Validation numérique : décimal >= 0."""
    dv = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0"
    )
    dv.error = "Saisissez une valeur numérique positive ou nulle."
    dv.errorTitle = "Validation numérique"
    ws.add_data_validation(dv)
    dv.add(plage)


def mise_en_page_paysage(ws):
    """Configure l'orientation paysage et l'ajustement à la largeur."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ----------------------------------------------------------------
# FEUILLE 1 : PARAMÈTRES
# ----------------------------------------------------------------

def creer_parametres(wb):
    ws = wb.create_sheet("Paramètres")
    style_titre(ws, "B2:D2", "PARAMÈTRES GÉNÉRAUX")

    champs = [
        ("Nom de l'entreprise", "Entreprise Exemple"),
        ("Site principal", "Site A"),
        ("Devise", "FCFA"),
        ("Prix du gasoil", 750),
        ("Facteur d'émission CO₂ gasoil", 2.68),
        ("Seuil d'alerte consommation L/h", 5),
        ("Seuil d'alerte charge faible %", 30),
        ("Seuil d'alerte fuite L", 1),
    ]

    for i, (libelle, valeur) in enumerate(champs, start=4):
        ws.cell(row=i, column=2, value=libelle)
        ws.cell(row=i, column=4, value=valeur)
        style_cellule(ws, f"B{i}", "normal")
        style_cellule(ws, f"D{i}", "saisie")

    definir_largeurs(ws, {"A": 3, "B": 40, "C": 3, "D": 25})

    wb.defined_names.add(DefinedName("PrixGasoil", attr_text="Paramètres!$D$7"))
    wb.defined_names.add(DefinedName("FacteurCO2", attr_text="Paramètres!$D$8"))

    return ws


# ----------------------------------------------------------------
# FEUILLE 2 : RÉFÉRENTIELS
# ----------------------------------------------------------------

def creer_referentiels(wb):
    ws = wb.create_sheet("Référentiels")
    style_titre(ws, "A1:H1", "RÉFÉRENTIELS — LISTES DE VALEURS")

    entetes = [
        "Sites", "Équipements", "Opérateurs", "Types de charge",
        "Débitmètres", "Unités eau", "Oui/Non", "Sources d'énergie"
    ]

    for col, titre in enumerate(entetes, start=1):
        ws.cell(row=3, column=col, value=titre)

    style_entetes(ws, 3, range(1, 9))

    donnees = {
        1: ["Site A", "Site B", "Site C"],
        2: ["GE-01", "GE-02", "GE-03"],
        3: ["Opérateur 1", "Opérateur 2", "Opérateur 3"],
        4: ["Pompage", "Éclairage", "Mixte", "Autre"],
        5: ["DN-01", "DN-02", "DN-03"],
        6: ["m³", "L"],
        7: ["Oui", "Non"],
        8: ["Gasoil", "Solaire", "Réseau", "Mixte"],
    }

    for col, valeurs in donnees.items():
        for i, valeur in enumerate(valeurs, start=4):
            c = ws.cell(row=i, column=col, value=valeur)
            style_cellule(ws, c.coordinate, "saisie", "center")

    definir_largeurs(ws, {get_column_letter(i): 20 for i in range(1, 9)})
    return ws


# ----------------------------------------------------------------
# FEUILLE 3 : SAISIE GASOIL
# ----------------------------------------------------------------

def creer_saisie_gasoil(wb):
    ws = wb.create_sheet("Saisie_Gasoil")
    style_titre(
        ws, "A1:Y1",
        "SUIVI DE LA CONSOMMATION DE GASOIL — GROUPE ÉLECTROGÈNE",
        14
    )

    ws["A2"] = (
        "Méthode : plein à plein — mesurer le niveau avant appoint "
        "et après consommation."
    )
    ws["A2"].font = Font(italic=True, color=BLEU_FONCE)

    entetes = [
        "N°", "Date", "Site", "Code équipement", "Type de charge",
        "Heure début", "Heure fin", "Compteur début (h)", "Compteur fin (h)",
        "Durée mesurée (h)", "Niveau initial référence", "Niveau final avant appoint",
        "Gasoil ajouté (L)", "Autres ajouts (L)", "Fuites / pertes (L)",
        "Conso corrigée (L)", "Conso moyenne (L/h)", "Charge estimée (%)",
        "Énergie produite (kWh)", "Conso spécifique (L/kWh)",
        "État / conditions", "Opérateur", "Observations / anomalies",
        "Vérif. niveau identique ?", "Validation responsable"
    ]

    for col, titre in enumerate(entetes, start=1):
        ws.cell(row=3, column=col, value=titre)

    style_entetes(ws, 3, range(1, 26))

    colonnes_saisie = list("BCDEFGHIKLMNORUVWXY")
    colonnes_calcul = list("AJPQST")

    for ligne in range(4, 104):
        ws[f"A{ligne}"] = f'=IF(B{ligne}="","",ROW()-3)'
        ws[f"J{ligne}"] = f'=IF(OR(H{ligne}="",I{ligne}=""),"",I{ligne}-H{ligne})'

        ws[f"P{ligne}"] = (
            f'=IF(M{ligne}="","",M{ligne}'
            f'+IF(N{ligne}="",0,N{ligne})'
            f'-IF(O{ligne}="",0,O{ligne}))'
        )

        ws[f"Q{ligne}"] = f'=IFERROR(P{ligne}/J{ligne},"")'
        ws[f"T{ligne}"] = f'=IFERROR(P{ligne}/S{ligne},"")'

        for col in colonnes_saisie:
            style_cellule(ws, f"{col}{ligne}", "saisie", "center")

        for col in colonnes_calcul:
            style_cellule(ws, f"{col}{ligne}", "calcul", "center")

    tab = Table(displayName="tblGasoil", ref="A3:Y103")
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False
    )
    ws.add_table(tab)

    ajouter_validation_liste(ws, "=Référentiels!$A$4:$A$53", "C4:C103")
    ajouter_validation_liste(ws, "=Référentiels!$B$4:$B$53", "D4:D103")
    ajouter_validation_liste(ws, "=Référentiels!$D$4:$D$53", "E4:E103")
    ajouter_validation_liste(ws, "=Référentiels!$C$4:$C$53", "V4:V103")
    ajouter_validation_liste(ws, "=Référentiels!$G$4:$G$53", "X4:X103")

    ajouter_validation_decimal_positif(ws, "M4:O103")

    ws.conditional_formatting.add(
        "O4:O103",
        CellIsRule(operator="greaterThan", formula=["0"], fill=fill(ROUGE))
    )

    ws.conditional_formatting.add(
        "X4:X103",
        FormulaRule(formula=['X4="Non"'], fill=fill(ORANGE))
    )

    ws.conditional_formatting.add(
        "X4:X103",
        FormulaRule(formula=['X4="Oui"'], fill=fill(VERT))
    )

    ws.freeze_panes = "C4"
    mise_en_page_paysage(ws)

    definir_largeurs(ws, {
        "A": 6, "B": 12, "C": 14, "D": 16, "E": 15,
        "F": 11, "G": 11, "H": 15, "I": 15, "J": 14,
        "K": 18, "L": 20, "M": 15, "N": 15, "O": 15,
        "P": 16, "Q": 16, "R": 15, "S": 17, "T": 18,
        "U": 18, "V": 16, "W": 28, "X": 20, "Y": 20
    })

    return ws


# ----------------------------------------------------------------
# FEUILLE 4 : SAISIE EAU
# ----------------------------------------------------------------

def creer_saisie_eau(wb):
    ws = wb.create_sheet("Saisie_Eau")
    style_titre(
        ws, "A1:N1",
        "SUIVI DE LA PRODUCTION D'EAU — DÉBITMÈTRES",
        14
    )

    entetes = [
        "N°", "Date", "Heure", "Site", "Débitmètre / conduite",
        "Index début", "Index fin", "Unité", "Volume", "Durée (h)",
        "Débit moyen (m³/h)", "Énergie utilisée (kWh)",
        "Source d'énergie", "Observations"
    ]

    for col, titre in enumerate(entetes, start=1):
        ws.cell(row=3, column=col, value=titre)

    style_entetes(ws, 3, range(1, 15))

    for ligne in range(4, 204):
        ws[f"A{ligne}"] = f'=IF(B{ligne}="","",ROW()-3)'
        ws[f"I{ligne}"] = f'=IF(OR(F{ligne}="",G{ligne}=""),"",G{ligne}-F{ligne})'
        ws[f"K{ligne}"] = f'=IFERROR(I{ligne}/J{ligne},"")'

        for col in "BCDEFGHJLMN":
            style_cellule(ws, f"{col}{ligne}", "saisie", "center")

        for col in "AIK":
            style_cellule(ws, f"{col}{ligne}", "calcul", "center")

    tab = Table(displayName="tblEau", ref="A3:N203")
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True
    )
    ws.add_table(tab)

    ajouter_validation_liste(ws, "=Référentiels!$A$4:$A$53", "D4:D203")
    ajouter_validation_liste(ws, "=Référentiels!$E$4:$E$53", "E4:E203")
    ajouter_validation_liste(ws, "=Référentiels!$F$4:$F$53", "H4:H203")
    ajouter_validation_liste(ws, "=Référentiels!$H$4:$H$53", "M4:M203")

    ws.freeze_panes = "C4"
    mise_en_page_paysage(ws)

    definir_largeurs(ws, {
        "A": 6, "B": 12, "C": 10, "D": 14, "E": 20,
        "F": 14, "G": 14, "H": 10, "I": 14, "J": 12,
        "K": 18, "L": 18, "M": 18, "N": 30
    })

    return ws


# ----------------------------------------------------------------
# FEUILLE 5 : SYNTHÈSE GASOIL
# ----------------------------------------------------------------

def creer_synthese_gasoil(wb):
    ws = wb.create_sheet("Synthèse_Gasoil")
    style_titre(ws, "A1:C1", "SYNTHÈSE AUTOMATIQUE — GASOIL", 14)

    kpis = [
        ("Nombre d'essais", '=COUNT(Saisie_Gasoil!B4:B103)'),
        ("Total heures", '=SUM(Saisie_Gasoil!J4:J103)'),
        ("Total gasoil corrigé", '=SUM(Saisie_Gasoil!P4:P103)'),
        (
            "Consommation moyenne pondérée",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Gasoil!J4:J103),"")'
        ),
        ("Énergie totale", '=SUM(Saisie_Gasoil!S4:S103)'),
        (
            "Consommation spécifique globale",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Gasoil!S4:S103),"")'
        ),
        ("Coût total gasoil", '=SUM(Saisie_Gasoil!P4:P103)*PrixGasoil'),
        ("Émissions CO₂ estimées", '=SUM(Saisie_Gasoil!P4:P103)*FacteurCO2'),
        ("Fuite totale constatée", '=SUM(Saisie_Gasoil!O4:O103)'),
    ]

    for i, (libelle, formule) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=libelle)
        ws.cell(row=i, column=3, value=formule)
        style_cellule(ws, f"A{i}", "normal")
        style_cellule(ws, f"C{i}", "calcul", "center")

    style_section(ws, "E3:F3", "DÉTAIL PAR SITE")

    ws["E4"] = "Site"
    ws["F4"] = "Gasoil total (L)"
    style_entetes(ws, 4, [5, 6])

    sites = ["Site A", "Site B", "Site C"]

    for i, site in enumerate(sites, start=5):
        ws.cell(row=i, column=5, value=site)
        ws.cell(
            row=i, column=6,
            value=f'=SUMIF(Saisie_Gasoil!C4:C103,E{i},Saisie_Gasoil!P4:P103)'
        )
        style_cellule(ws, f"E{i}", "saisie", "center")
        style_cellule(ws, f"F{i}", "calcul", "center")

    definir_largeurs(ws, {
        "A": 34, "B": 3, "C": 22, "D": 3, "E": 18, "F": 22
    })

    return ws


# ----------------------------------------------------------------
# FEUILLE 6 : SYNTHÈSE EAU
# ----------------------------------------------------------------

def creer_synthese_eau(wb):
    ws = wb.create_sheet("Synthèse_Eau")
    style_titre(
        ws, "A1:C1",
        "SYNTHÈSE AUTOMATIQUE — PRODUCTION D'EAU",
        14
    )

    kpis = [
        ("Nombre de relevés", '=COUNT(Saisie_Eau!B4:B203)'),
        ("Volume total", '=SUM(Saisie_Eau!I4:I203)'),
        ("Durée totale", '=SUM(Saisie_Eau!J4:J203)'),
        (
            "Débit moyen global",
            '=IFERROR(SUM(Saisie_Eau!I4:I203)/SUM(Saisie_Eau!J4:J203),"")'
        ),
        ("Énergie totale", '=SUM(Saisie_Eau!L4:L203)'),
        (
            "Ratio gasoil/eau (L/m³)",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Eau!I4:I203),"")'
        ),
        (
            "Ratio kWh/m³",
            '=IFERROR(SUM(Saisie_Eau!L4:L203)/SUM(Saisie_Eau!I4:I203),"")'
        ),
    ]

    for i, (libelle, formule) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=libelle)
        ws.cell(row=i, column=3, value=formule)
        style_cellule(ws, f"A{i}", "normal")
        style_cellule(ws, f"C{i}", "calcul", "center")

    style_section(ws, "E3:F3", "DÉTAIL PAR SITE")

    ws["E4"] = "Site"
    ws["F4"] = "Volume total"
    style_entetes(ws, 4, [5, 6])

    sites = ["Site A", "Site B", "Site C"]

    for i, site in enumerate(sites, start=5):
        ws.cell(row=i, column=5, value=site)
        ws.cell(
            row=i, column=6,
            value=f'=SUMIF(Saisie_Eau!D4:D203,E{i},Saisie_Eau!I4:I203)'
        )
        style_cellule(ws, f"E{i}", "saisie", "center")
        style_cellule(ws, f"F{i}", "calcul", "center")

    definir_largeurs(ws, {
        "A": 32, "B": 3, "C": 22, "D": 3, "E": 18, "F": 22
    })

    return ws


# ----------------------------------------------------------------
# FEUILLE 7 : TABLEAU DE BORD
# ----------------------------------------------------------------

def creer_tableau_bord(wb):
    ws = wb.create_sheet("Tableau_Bord")
    style_titre(ws, "A1:H1", "TABLEAU DE BORD — GASOIL & EAU", 16)

    cartes = [
        ("Gasoil total (L)", '=SUM(Saisie_Gasoil!P4:P103)'),
        (
            "Conso moyenne (L/h)",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Gasoil!J4:J103),"")'
        ),
        (
            "Conso spécifique (L/kWh)",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Gasoil!S4:S103),"")'
        ),
        ("Coût total", '=SUM(Saisie_Gasoil!P4:P103)*PrixGasoil'),
        ("Volume eau", '=SUM(Saisie_Eau!I4:I203)'),
        (
            "Débit moyen",
            '=IFERROR(SUM(Saisie_Eau!I4:I203)/SUM(Saisie_Eau!J4:J203),"")'
        ),
        (
            "L gasoil/m³ eau",
            '=IFERROR(SUM(Saisie_Gasoil!P4:P103)/SUM(Saisie_Eau!I4:I203),"")'
        ),
        ("CO₂ émis", '=SUM(Saisie_Gasoil!P4:P103)*FacteurCO2'),
    ]

    colonnes = ["A", "B", "C", "D", "E", "F", "G", "H"]

    for i, (titre, formule) in enumerate(cartes):
        col = colonnes[i]

        ws[f"{col}3"] = titre
        ws[f"{col}4"] = formule

        ws[f"{col}3"].font = Font(bold=True, color="FFFFFF", size=10)
        ws[f"{col}3"].fill = fill(BLEU_MOYEN)
        ws[f"{col}3"].alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        ws[f"{col}4"].font = Font(bold=True, size=16, color=BLEU_FONCE)
        ws[f"{col}4"].fill = fill(BLEU_CLAIR)
        ws[f"{col}4"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        for coord in [f"{col}3", f"{col}4"]:
            ws[coord].border = bordure_fine()

    ws.row_dimensions[3].height = 34
    ws.row_dimensions[4].height = 40

    chart_gasoil = BarChart()
    chart_gasoil.title = "Gasoil total par site (L)"
    chart_gasoil.height = 9
    chart_gasoil.width = 16

    data = Reference(
        wb["Synthèse_Gasoil"],
        min_col=6,
        min_row=4,
        max_row=7
    )
    cats = Reference(
        wb["Synthèse_Gasoil"],
        min_col=5,
        min_row=5,
        max_row=7
    )

    chart_gasoil.add_data(data, titles_from_data=True)
    chart_gasoil.set_categories(cats)
    ws.add_chart(chart_gasoil, "A10")

    chart_eau = BarChart()
    chart_eau.title = "Volume d'eau par site (m³)"
    chart_eau.height = 9
    chart_eau.width = 16

    data = Reference(
        wb["Synthèse_Eau"],
        min_col=6,
        min_row=4,
        max_row=7
    )
    cats = Reference(
        wb["Synthèse_Eau"],
        min_col=5,
        min_row=5,
        max_row=7
    )

    chart_eau.add_data(data, titles_from_data=True)
    chart_eau.set_categories(cats)
    ws.add_chart(chart_eau, "E10")

    definir_largeurs(ws, {col: 18 for col in colonnes})
    return ws


# ----------------------------------------------------------------
# FEUILLE 8 : FICHE TERRAIN
# ----------------------------------------------------------------

def creer_fiche_terrain(wb):
    ws = wb.create_sheet("Fiche_Terrain")
    style_titre(
        ws, "A1:D1",
        "FICHE TERRAIN — MESURE PLEIN À PLEIN",
        14
    )

    champs = [
        ("Site", ""),
        ("Groupe / code", ""),
        ("Date", ""),
        ("Opérateur", ""),
        ("Puissance kVA", ""),
        ("Charge / équipement", ""),
        ("Compteur début", ""),
        ("Compteur fin", ""),
    ]

    for i, (libelle, valeur) in enumerate(champs, start=3):
        ws.cell(row=i, column=1, value=libelle)
        ws.cell(row=i, column=2, value=valeur)
        style_cellule(ws, f"A{i}", "normal")
        style_cellule(ws, f"B{i}", "saisie")

    ws["A12"] = "Durée (h)"
    ws["B12"] = '=IF(OR(B9="",B10=""),"",B10-B9)'

    ws["A13"] = "Consommation estimée (L)"
    ws["B13"] = '=IF(B12="","",B12*5)'

    ws["A14"] = "Consommation moyenne (L/h)"
    ws["B14"] = '=IFERROR(B13/B12,"")'

    for ligne in range(12, 15):
        style_cellule(ws, f"A{ligne}", "normal")
        style_cellule(ws, f"B{ligne}", "calcul", "center")

    style_section(
        ws, "A16:D16",
        "SUIVI DE LA PRODUCTION D'EAU — DÉBITMÈTRES"
    )

    entetes_eau = [
        "Débitmètre", "Index début", "Index fin",
        "Volume", "Durée (h)", "Débit moyen"
    ]

    for col, titre in enumerate(entetes_eau, start=1):
        ws.cell(row=17, column=col, value=titre)

    style_entetes(ws, 17, range(1, 7))

    for ligne in range(18, 23):
        ws[f"D{ligne}"] = f'=IF(OR(B{ligne}="",C{ligne}=""),"",C{ligne}-B{ligne})'
        ws[f"F{ligne}"] = f'=IFERROR(D{ligne}/E{ligne},"")'

        for col in "ABCE":
            style_cellule(ws, f"{col}{ligne}", "saisie", "center")

        for col in "DF":
            style_cellule(ws, f"{col}{ligne}", "calcul", "center")

    style_section(ws, "A25:D25", "OBSERVATIONS")

    ws.merge_cells("A26:D29")
    ws["A26"] = ""
    style_cellule(ws, "A26", "saisie")

    ws["A31"] = "Signature opérateur :"
    ws["C31"] = "Signature responsable :"

    for coord in ["A31", "C31"]:
        ws[coord].font = Font(bold=True, color=BLEU_FONCE)

    mise_en_page_paysage(ws)

    definir_largeurs(ws, {
        "A": 28, "B": 22, "C": 22, "D": 22,
        "E": 14, "F": 16
    })

    return ws


# ----------------------------------------------------------------
# FEUILLE 9 : MODE D'EMPLOI
# ----------------------------------------------------------------

def creer_mode_emploi(wb):
    ws = wb.create_sheet("Mode_Emploi")
    style_titre(ws, "A1:B1", "MODE D'EMPLOI", 16)

    etapes = [
        "1. Renseignez d'abord les paramètres généraux : entreprise, prix du gasoil, facteur CO₂ et seuils d'alerte.",
        "2. Complétez les référentiels : sites, groupes électrogènes, opérateurs, débitmètres et sources d'énergie.",
        "3. Dans Saisie_Gasoil, enregistrez chaque essai selon la méthode plein à plein.",
        "4. Saisissez le compteur début et fin, puis les volumes ajoutés et les fuites éventuelles.",
        "5. Les colonnes grisées calculent automatiquement la durée, la consommation et les ratios.",
        "6. Dans Saisie_Eau, saisissez les index de débitmètres, la durée et l'énergie utilisée.",
        "7. Le volume et le débit moyen sont calculés automatiquement.",
        "8. Consultez les synthèses et le tableau de bord pour suivre les performances par site.",
        "9. Utilisez les macros VBA, si elles sont activées, pour ajouter des lignes, naviguer et exporter en PDF.",
        "10. Sauvegardez une copie mensuelle du classeur afin d'assurer la traçabilité historique des données."
    ]

    for i, texte in enumerate(etapes, start=3):
        ws.cell(row=i, column=1, value=texte)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        style_cellule(ws, f"A{i}", "normal")
        ws.row_dimensions[i].height = 30

    definir_largeurs(ws, {"A": 100, "B": 20})
    return ws


# ----------------------------------------------------------------
# GÉNÉRATION DES MODULES VBA
# ----------------------------------------------------------------

def generer_modules_vba():
    """Crée le dossier vba/ et écrit les modules VBA."""
    os.makedirs(DOSSIER_VBA, exist_ok=True)

    mod_essais = '''Attribute VB_Name = "modEssais"
Option Explicit

Public Sub AjouterEssaiGasoil()
    Dim tbl As ListObject
    Dim nouvelleLigne As ListRow

    Set tbl = ThisWorkbook.Worksheets("Saisie_Gasoil").ListObjects("tblGasoil")
    Set nouvelleLigne = tbl.ListRows.Add

    nouvelleLigne.Range.Cells(1, 2).Value = Date
    nouvelleLigne.Range.Cells(1, 24).Value = "Oui"

    MsgBox "Nouvelle ligne ajoutée dans tblGasoil.", vbInformation, "Gasoil"
End Sub

Public Sub AjouterReleveEau()
    Dim tbl As ListObject
    Dim nouvelleLigne As ListRow

    Set tbl = ThisWorkbook.Worksheets("Saisie_Eau").ListObjects("tblEau")
    Set nouvelleLigne = tbl.ListRows.Add

    nouvelleLigne.Range.Cells(1, 2).Value = Date

    MsgBox "Nouveau relevé ajouté dans tblEau.", vbInformation, "Production d'eau"
End Sub

Public Sub RazFicheTerrain()
    Dim reponse As VbMsgBoxResult

    reponse = MsgBox("Voulez-vous réinitialiser la fiche terrain ?", vbYesNo + vbQuestion, "Confirmation")

    If reponse <> vbYes Then Exit Sub

    With ThisWorkbook.Worksheets("Fiche_Terrain")
        .Range("B3:B11").ClearContents
        .Range("A18:F22").ClearContents
        .Range("A26:D29").ClearContents
    End With

    MsgBox "Fiche terrain réinitialisée.", vbInformation, "Terminé"
End Sub
'''

    mod_export = '''Attribute VB_Name = "modExport"
Option Explicit

Public Sub ExporterTableauBordPDF()
    Dim chemin As String
    Dim nomFichier As String

    nomFichier = "Tableau_Bord_" & Format(Date, "yyyy-mm-dd") & ".pdf"
    chemin = ThisWorkbook.Path & Application.PathSeparator & nomFichier

    ThisWorkbook.Worksheets("Tableau_Bord").ExportAsFixedFormat _
        Type:=xlTypePDF, _
        Filename:=chemin, _
        Quality:=xlQualityStandard, _
        OpenAfterPublish:=True

    MsgBox "PDF créé : " & chemin, vbInformation, "Export PDF"
End Sub

Public Sub ExporterFicheTerrainPDF()
    Dim chemin As String
    Dim nomFichier As String

    nomFichier = "Fiche_Terrain_" & Format(Date, "yyyy-mm-dd") & ".pdf"
    chemin = ThisWorkbook.Path & Application.PathSeparator & nomFichier

    ThisWorkbook.Worksheets("Fiche_Terrain").ExportAsFixedFormat _
        Type:=xlTypePDF, _
        Filename:=chemin, _
        Quality:=xlQualityStandard, _
        OpenAfterPublish:=True

    MsgBox "PDF créé : " & chemin, vbInformation, "Export PDF"
End Sub

Public Sub ActualiserTout()
    Application.CalculateFull
    MsgBox "Tous les calculs ont été actualisés.", vbInformation, "Actualisation"
End Sub
'''

    mod_navigation = '''Attribute VB_Name = "modNavigation"
Option Explicit

Public Sub AllerTableauBord()
    ThisWorkbook.Worksheets("Tableau_Bord").Activate
End Sub

Public Sub AllerSaisieGasoil()
    ThisWorkbook.Worksheets("Saisie_Gasoil").Activate
End Sub

Public Sub AllerSaisieEau()
    ThisWorkbook.Worksheets("Saisie_Eau").Activate
End Sub

Public Sub AllerFicheTerrain()
    ThisWorkbook.Worksheets("Fiche_Terrain").Activate
End Sub
'''

    this_workbook = '''Attribute VB_Name = "ThisWorkbook"
Option Explicit

Private Sub Workbook_Open()
    ThisWorkbook.Worksheets("Tableau_Bord").Activate
End Sub
'''

    fichiers = {
        "modEssais.bas": mod_essais,
        "modExport.bas": mod_export,
        "modNavigation.bas": mod_navigation,
        "ThisWorkbook.cls": this_workbook,
    }

    for nom, contenu in fichiers.items():
        chemin = os.path.join(DOSSIER_VBA, nom)
        with open(chemin, "w", encoding="utf-8") as f:
            f.write(contenu)

    return list(fichiers.keys())


# ----------------------------------------------------------------
# INJECTION VBA ET CONVERSION EN XLSM
# ----------------------------------------------------------------

def injecter_macros_vba(chemin_xlsx, chemin_xlsm, fichiers_vba):
    """
    Ouvre le fichier .xlsx, importe les modules VBA,
    puis enregistre le classeur en .xlsm.

    FileFormat = 52 correspond à :
    xlOpenXMLWorkbookMacroEnabled (.xlsm)
    """
    try:
        import win32com.client
    except ImportError:
        return False

    excel = None
    classeur = None

    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        classeur = excel.Workbooks.Open(os.path.abspath(chemin_xlsx))

        for nom_fichier in fichiers_vba:
            chemin_module = os.path.abspath(os.path.join(DOSSIER_VBA, nom_fichier))
            classeur.VBProject.VBComponents.Import(chemin_module)

        classeur.SaveAs(
            os.path.abspath(chemin_xlsm),
            FileFormat=52
        )

        return True

    except Exception as e:
        print(f"[AVERTISSEMENT] Injection VBA impossible : {e}")
        return False

    finally:
        try:
            if classeur is not None:
                classeur.Close(SaveChanges=False)
            if excel is not None:
                excel.Quit()
        except Exception:
            pass


# ----------------------------------------------------------------
# PROGRAMME PRINCIPAL
# ----------------------------------------------------------------

def main():
    try:
        print("Création du classeur Excel .xlsx...")

        wb = Workbook()
        wb.remove(wb.active)

        creer_parametres(wb)
        creer_referentiels(wb)
        creer_saisie_gasoil(wb)
        creer_saisie_eau(wb)
        creer_synthese_gasoil(wb)
        creer_synthese_eau(wb)
        creer_tableau_bord(wb)
        creer_fiche_terrain(wb)
        creer_mode_emploi(wb)

        wb.save(FICHIER_XLSX)

        chemin_xlsx = os.path.abspath(FICHIER_XLSX)
        chemin_xlsm = os.path.abspath(FICHIER_XLSM)

        print(f"Classeur .xlsx créé : {chemin_xlsx}")

        print("Génération des modules VBA...")
        fichiers_vba = generer_modules_vba()

        print("Tentative d'injection automatique des macros VBA...")
        injection_ok = injecter_macros_vba(
            chemin_xlsx,
            chemin_xlsm,
            fichiers_vba
        )

        if injection_ok:
            print("\nMacros VBA injectées avec succès.")
            print(f"Classeur final .xlsm créé : {chemin_xlsm}")
        else:
            print("\n" + "=" * 70)
            print("IMPORT MANUEL DES MACROS VBA")
            print("=" * 70)
            print("Le fichier .xlsx a été créé, mais les macros n'ont pas été injectées.")
            print("\nProcédure :")
            print("1. Ouvrez le fichier .xlsx dans Excel.")
            print("2. Appuyez sur Alt + F11 pour ouvrir l'éditeur VBA.")
            print("3. Faites : Fichier > Importer un fichier...")
            print("4. Importez successivement les fichiers suivants :")

            for fichier in fichiers_vba:
                print(f"   - {os.path.abspath(os.path.join(DOSSIER_VBA, fichier))}")

            print("5. Enregistrez le classeur au format :")
            print("   Classeur Excel prenant en charge les macros (*.xlsm)")
            print("=" * 70)

        return chemin_xlsm if injection_ok else chemin_xlsx

    except PermissionError:
        print("[ERREUR] Le fichier Excel est probablement ouvert.")
        print("Fermez-le puis relancez le script.")
        sys.exit(1)

    except Exception as e:
        print(f"[ERREUR] Une erreur est survenue : {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()